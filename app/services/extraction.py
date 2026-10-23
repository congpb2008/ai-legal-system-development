from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.services.parser import SourceLine

logger = logging.getLogger(__name__)


class ExtractionError(RuntimeError):
    pass


class OcrUnavailableError(ExtractionError):
    pass


@dataclass
class ExtractedPage:
    page_number: int
    width: float | None
    height: float | None
    text: str
    lines: list[SourceLine] = field(default_factory=list)
    text_source: str = "text"
    ocr_confidence: float | None = None
    image_path: str | None = None


_paddle_engine: Any = None


def _get_paddle(lang: str):
    global _paddle_engine
    if _paddle_engine is None:
        try:
            from paddleocr import PaddleOCR  # type: ignore
        except Exception as exc:  # pragma: no cover - optional dependency
            raise OcrUnavailableError(
                "PaddleOCR is not installed. Use `uv sync --extra ocr` or the Docker image."
            ) from exc
        _paddle_engine = PaddleOCR(use_angle_cls=True, lang=lang, show_log=False)
    return _paddle_engine


def _ocr_with_paddle(image, lang: str) -> tuple[list[SourceLine], float]:
    import numpy as np

    engine = _get_paddle(lang)
    result = engine.ocr(np.array(image), cls=True)
    raw_lines: list[SourceLine] = []
    scores: list[float] = []
    for page_result in result or []:
        for item in page_result or []:
            box, (text, score) = item[0], item[1]
            if not text or not str(text).strip():
                continue
            xs = [point[0] for point in box]
            ys = [point[1] for point in box]
            raw_lines.append(
                SourceLine(
                    text=str(text),
                    page_number=0,
                    bbox=[float(min(xs)), float(min(ys)), float(max(xs)), float(max(ys))],
                )
            )
            scores.append(float(score))
    raw_lines.sort(key=lambda line: ((line.bbox or [0, 0, 0, 0])[1], (line.bbox or [0, 0, 0, 0])[0]))
    confidence = sum(scores) / len(scores) if scores else 0.0
    return raw_lines, confidence


_TESSERACT_LANG = {"vi": "vie", "en": "eng", "zh": "chi_sim", "ja": "jpn", "ko": "kor"}


def _scale_ocr_bboxes(
    lines: list[SourceLine],
    image_width: int,
    image_height: int,
    page_width: float,
    page_height: float,
) -> None:
    """Convert pixel coordinates from OCR raster images back to PDF points."""
    if not image_width or not image_height:
        return
    scale_x = page_width / image_width
    scale_y = page_height / image_height
    for line in lines:
        if not line.bbox or len(line.bbox) != 4:
            continue
        line.bbox = [
            line.bbox[0] * scale_x,
            line.bbox[1] * scale_y,
            line.bbox[2] * scale_x,
            line.bbox[3] * scale_y,
        ]


def _ocr_with_tesseract(image, lang: str) -> tuple[list[SourceLine], float]:
    try:
        import pytesseract  # type: ignore
        from pytesseract import Output  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency
        raise OcrUnavailableError(
            "Tesseract OCR is not installed. Install tesseract-ocr + tesseract-ocr-vie and pytesseract."
        ) from exc

    tesseract_lang = _TESSERACT_LANG.get(lang, lang)
    data = pytesseract.image_to_data(image, lang=tesseract_lang, output_type=Output.DICT)
    grouped: dict[tuple[int, int, int, int], dict[str, Any]] = {}
    for index, text in enumerate(data.get("text", [])):
        if not str(text).strip():
            continue
        key = (data["page_num"][index], data["block_num"][index], data["par_num"][index], data["line_num"][index])
        entry = grouped.setdefault(key, {"words": [], "left": [], "top": [], "right": [], "bottom": [], "conf": []})
        entry["words"].append(str(text))
        left, top = data["left"][index], data["top"][index]
        width, height = data["width"][index], data["height"][index]
        entry["left"].append(left)
        entry["top"].append(top)
        entry["right"].append(left + width)
        entry["bottom"].append(top + height)
        try:
            entry["conf"].append(float(data["conf"][index]))
        except (TypeError, ValueError):
            pass

    lines: list[SourceLine] = []
    confidences: list[float] = []
    for entry in grouped.values():
        text = " ".join(entry["words"]).strip()
        if not text:
            continue
        lines.append(
            SourceLine(
                text=text,
                page_number=0,
                bbox=[
                    float(min(entry["left"])),
                    float(min(entry["top"])),
                    float(max(entry["right"])),
                    float(max(entry["bottom"])),
                ],
            )
        )
        confidences.extend(entry["conf"])
    lines.sort(key=lambda line: ((line.bbox or [0, 0, 0, 0])[1], (line.bbox or [0, 0, 0, 0])[0]))
    confidence = sum(confidences) / len(confidences) / 100.0 if confidences else 0.0
    return lines, confidence


def _lines_from_pdf_page(page) -> list[SourceLine]:
    result = page.get_text("dict")
    lines: list[SourceLine] = []
    for block in result.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            text = "".join(span.get("text", "") for span in spans).strip()
            if not text:
                continue
            bbox = line.get("bbox")
            font_size = max((float(span.get("size", 0)) for span in spans), default=0.0)
            bold = any(
                "bold" in str(span.get("font", "")).lower() or (int(span.get("flags", 0)) & 16)
                for span in spans
            )
            lines.append(
                SourceLine(
                    text=text,
                    page_number=0,
                    bbox=[float(value) for value in bbox] if bbox else None,
                    font_size=font_size,
                    bold=bool(bold),
                )
            )
    lines.sort(key=lambda line: ((line.bbox or [0, 0, 0, 0])[1], (line.bbox or [0, 0, 0, 0])[0]))
    return lines


def extract_pdf(
    path: Path,
    *,
    ocr_provider: str = "auto",
    ocr_lang: str = "vi",
    ocr_dpi: int = 220,
    min_text_chars: int = 40,
    output_dir: Path | None = None,
) -> list[ExtractedPage]:
    try:
        import fitz  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise ExtractionError("PyMuPDF is required for PDF extraction.") from exc

    document = fitz.open(path)
    pages: list[ExtractedPage] = []
    try:
        for page_index, page in enumerate(document, start=1):
            lines = _lines_from_pdf_page(page)
            page_text = "\n".join(line.text for line in lines).strip()
            source = "text"
            confidence: float | None = 1.0
            image_path: str | None = None

            if len(page_text) < min_text_chars:
                if ocr_provider == "none":
                    if not page_text:
                        raise OcrUnavailableError(
                            f"Page {page_index} has no text layer and OCR is disabled."
                        )
                else:
                    pixmap = page.get_pixmap(dpi=ocr_dpi, alpha=False)
                    from PIL import Image
                    import io

                    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                    try:
                        if ocr_provider in {"auto", "paddle"}:
                            try:
                                ocr_lines, confidence = _ocr_with_paddle(image, ocr_lang)
                            except OcrUnavailableError:
                                if ocr_provider == "paddle":
                                    raise
                                ocr_lines, confidence = _ocr_with_tesseract(image, ocr_lang)
                            except Exception:
                                if ocr_provider == "paddle":
                                    raise
                                logger.exception("PaddleOCR failed; falling back to Tesseract")
                                ocr_lines, confidence = _ocr_with_tesseract(image, ocr_lang)
                        elif ocr_provider == "tesseract":
                            ocr_lines, confidence = _ocr_with_tesseract(image, ocr_lang)
                        else:
                            raise ExtractionError(f"Unknown OCR provider: {ocr_provider}")

                        if ocr_lines:
                            _scale_ocr_bboxes(
                                ocr_lines,
                                image.width,
                                image.height,
                                float(page.rect.width),
                                float(page.rect.height),
                            )
                        for line in ocr_lines:
                            line.page_number = page_index
                        if ocr_lines:
                            lines = ocr_lines
                            page_text = "\n".join(line.text for line in lines).strip()
                            source = "ocr"
                            if output_dir is not None:
                                output_dir.mkdir(parents=True, exist_ok=True)
                                image_file = output_dir / f"page-{page_index:04d}.jpg"
                                image.convert("RGB").save(image_file, quality=85)
                                image_path = str(image_file)
                    except OcrUnavailableError:
                        if not page_text:
                            raise
                        # A watermark or partial text layer is still better than nothing.
                        source = "text-low"
                        confidence = None

            for line in lines:
                line.page_number = page_index
            pages.append(
                ExtractedPage(
                    page_number=page_index,
                    width=float(page.rect.width),
                    height=float(page.rect.height),
                    text=page_text,
                    lines=lines,
                    text_source=source,
                    ocr_confidence=confidence if source == "ocr" else None,
                    image_path=image_path,
                )
            )
    finally:
        document.close()
    return pages


def extract_docx(path: Path) -> list[ExtractedPage]:
    try:
        from docx import Document as DocxDocument  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise ExtractionError("python-docx is required for DOCX extraction.") from exc

    doc = DocxDocument(str(path))
    lines: list[SourceLine] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name if paragraph.style else "").lower()
        bold = "heading" in style_name or "title" in style_name
        lines.append(SourceLine(text=text, page_number=1, font_size=0.0, bold=bold))
    for table in doc.tables:
        for row in table.rows:
            text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if text:
                lines.append(SourceLine(text=text, page_number=1))
    text = "\n".join(line.text for line in lines).strip()
    return [
        ExtractedPage(
            page_number=1,
            width=None,
            height=None,
            text=text,
            lines=lines,
            text_source="docx",
            ocr_confidence=None,
        )
    ]


def extract_document(
    path: Path,
    *,
    mime_type: str = "application/pdf",
    original_filename: str = "",
    ocr_provider: str = "auto",
    ocr_lang: str = "vi",
    ocr_dpi: int = 220,
    min_text_chars: int = 40,
    output_dir: Path | None = None,
) -> list[ExtractedPage]:
    suffix = Path(original_filename or path.name).suffix.lower()
    if suffix == ".docx" or mime_type in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }:
        return extract_docx(path)
    if suffix == ".pdf" or mime_type == "application/pdf":
        return extract_pdf(
            path,
            ocr_provider=ocr_provider,
            ocr_lang=ocr_lang,
            ocr_dpi=ocr_dpi,
            min_text_chars=min_text_chars,
            output_dir=output_dir,
        )
    raise ExtractionError(f"Unsupported file type: {suffix or mime_type}")
